import json,re,sys
from pathlib import Path
from datetime import datetime,timezone
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_BREAK,WD_TAB_ALIGNMENT,WD_TAB_LEADER
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.part import Part
from docx.opc.packuri import PackURI
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from research_tables import tables
A=Path(__file__).resolve().parent
R=A.parents[1]
FIGURES={1:'architecture',2:'authorization-flow',3:'compromise-recovery',4:'quantum-migration'}
C=json.loads((A/'paper_content.json').read_text());T=tables(C)
pages=json.loads((A/'toc-pages.json').read_text()) if (A/'toc-pages.json').exists() else {}
doc=Document();section=doc.sections[0]
section.page_width=Inches(8.5);section.page_height=Inches(11)
section.top_margin=Inches(.78);section.bottom_margin=Inches(.74)
section.left_margin=Inches(.9);section.right_margin=Inches(.9)
section.header_distance=Inches(.3);section.footer_distance=Inches(.3)
section.different_first_page_header_footer=True
def font(style,name,size,bold=False):
 s=doc.styles[style];s.font.name=name;s.font.size=Pt(size);s.font.bold=bold;s.font.color.rgb=RGBColor(0,0,0)
 rf=s.element.get_or_add_rPr().rFonts
 for a in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
  if qn('w:'+a) in rf.attrib:del rf.attrib[qn('w:'+a)]
 s.paragraph_format.widow_control=True
 return s
normal=font('Normal','Times New Roman',11)
normal.paragraph_format.line_spacing=1.08;normal.paragraph_format.space_after=Pt(6)
for name,size in [('Title',28),('Subtitle',15),('Heading 1',14),('Heading 2',11.5)]:
 s=font(name,'Times New Roman',size,name!='Subtitle');s.paragraph_format.space_after=Pt(8)
 s.paragraph_format.space_before=Pt(13 if name.startswith('Heading') else 0)
 s.paragraph_format.keep_with_next=True
for name,size in [('Caption',9.5),('Reference',9.5),('Small',9.5),('TOC 1',10),('TOC 2',9.5)]:
 if name not in doc.styles:doc.styles.add_style(name,1)
 s=font(name,'Times New Roman',size);s.paragraph_format.space_after=Pt(5);s.paragraph_format.line_spacing=1.02
doc.styles['TOC 1'].paragraph_format.space_after=Pt(4)
doc.styles['TOC 2'].paragraph_format.left_indent=Inches(.2)
doc.styles['TOC 2'].paragraph_format.space_after=Pt(3)
for name in ['TOC 1','TOC 2']:
 doc.styles[name].paragraph_format.tab_stops.add_tab_stop(Inches(6.65),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
doc.core_properties.title=C['title'];doc.core_properties.subject=C['subtitle'];doc.core_properties.author=C['author']
doc.core_properties.created=datetime(2026,10,21,tzinfo=timezone.utc)
doc.core_properties.keywords='custody, threshold ECDSA, post quantum, transaction authorization, recovery'
h=section.header.paragraphs[0];h.text='INSTITUTIONAL DIGITAL ASSET CUSTODY';h.style='Small';h.runs[0].font.size=Pt(8)
f=section.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=f.add_run('Dimitris Konstantinou  |  ');r.font.size=Pt(8)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
def link(p,text,url=None,anchor=None):
 h=OxmlElement('w:hyperlink')
 if url:h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
 if anchor:h.set(qn('w:anchor'),anchor)
 r=OxmlElement('w:r');rp=OxmlElement('w:rPr');col=OxmlElement('w:color');col.set(qn('w:val'),'173D58');rp.append(col);r.append(rp)
 t=OxmlElement('w:t');t.text=text;r.append(t);h.append(r);p._p.append(h)
def para(t,style=None):return doc.add_paragraph(t,style)
headings=[]
def heading(text,level=1):
 p=doc.add_heading(text,level);anchor=f'h_{len(headings)+1:03d}'
 start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),str(len(headings)+1));start.set(qn('w:name'),anchor)
 end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(len(headings)+1));p._p.insert(0,start);p._p.append(end)
 headings.append({'text':text,'level':level,'anchor':anchor});return p
p=para('', 'Title');p.paragraph_format.space_before=Pt(76)
p.add_run('Securing Institutional').add_break();p.add_run('Digital Asset Custody')
p=para(C['subtitle'],'Subtitle');p.paragraph_format.space_after=Pt(28)
para(C['author']);p=para('Publication date  '+C['date']);p.paragraph_format.space_after=Pt(18)
p=para('GitHub  ');link(p,'github.com/demis1997','https://github.com/demis1997')
p=para('Research repository  ');link(p,'github.com/demis1997/quorum-custody','https://github.com/demis1997/quorum-custody')
p=para('Transaction authority, distributed state, recovery, and cryptographic migration');p.paragraph_format.space_before=Pt(26)
doc.add_page_break()
heading('Abstract');para(C['abstract'])
para('Keywords: institutional custody, threshold ECDSA, transaction authorization, distributed state, recovery, post quantum cryptography.','Small')
doc.add_page_break();para('Contents','Heading 1')
toc_anchor=doc.add_paragraph() # replaced after headings are known
doc.add_page_break()

def mr(t):
 r=OxmlElement('m:r');tx=OxmlElement('m:t');tx.text=t;r.append(tx);return r
def indexed(base,index,tag):
 e=OxmlElement('m:'+tag);b=OxmlElement('m:e');b.append(mr(base));i=OxmlElement('m:'+('sub' if tag=='sSub' else 'sup'));i.append(mr(index));e.extend([b,i]);return e
def equation(n):
 p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(10);p.paragraph_format.keep_with_next=True
 o=OxmlElement('m:oMath')
 if n==1:o.extend([mr('Q = xG,   R = kG,   r = '),indexed('R','x','sSub'),mr(' mod n,   s = '),indexed('k','−1','sSup'),mr('(z + rx) mod n')])
 else:o.extend([mr('C = v + gf,     '),indexed('B','reserved','sSub'),mr(' + C ≤ '),indexed('B','limit','sSub')])
 p._p.append(o)
captions={1:'Figure 1. Institutional custody trust boundaries. Dashed signer domains and external protection describe proposed institutional administration. The prototype uses one host account; no HSM or KMS isolation is claimed. Representative paths are shown; each signer needs independently governed recovery.',2:'Figure 2. Transaction authorization lifecycle. The dashed path marks submission by another holder after ordinary signed bytes exist. Application expiry constrains honest components; network transaction validation has no knowledge of private approval expiry.',3:'Figure 3. Compromise and recovery boundaries. Effects depend on the declared attacker capability, approval evidence, available signed bytes, and control of comparison checkpoints. A single compromised signer is not a compromised quorum.',4:'Figure 4. Proposed post quantum custody migration. The transition requires asset movement and protection of upgrade and recovery authority. These phases are proposed and have not been deployed or tested as a contract account.'}
def figure(n):
 p=doc.add_paragraph();p.paragraph_format.keep_with_next=True
 run=p.add_run();run.add_picture(str(R/'figures'/f'{FIGURES[n]}.png'),width=Inches(6.65))
 run._r.xpath('.//wp:docPr')[0].set('descr',captions[n])
 part=Part(PackURI(f'/word/media/figure-{n}.svg'),'image/svg+xml',(R/'figures'/f'{FIGURES[n]}.svg').read_bytes(),doc.part.package)
 rid=doc.part.relate_to(part,RT.IMAGE)
 blip=run._r.xpath('.//a:blip')[0]
 extlist=OxmlElement('a:extLst');ext=OxmlElement('a:ext');ext.set('uri','{96DAC541-7B7A-43D3-8B79-37D633B846F1}')
 from lxml import etree
 svg=etree.Element('{http://schemas.microsoft.com/office/drawing/2016/SVG/main}svgBlip',nsmap={'asvg':'http://schemas.microsoft.com/office/drawing/2016/SVG/main'});svg.set(qn('r:embed'),rid)
 ext.append(svg);extlist.append(ext);blip.append(extlist)
 para(captions[n],'Caption')
def table(name):
 caption,heads,rows=T[name]
 cap=para(caption,'Caption');cap.paragraph_format.keep_with_next=True
 t=doc.add_table(rows=1,cols=len(heads));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 widths={'threat':[1.0,1.5,1.45,1.85,.9],'experiments':[1.15,2.55,3.0],'sizes':[1.9,2.15,1.25,1.4],'benchmark':[3.1,1.8,1.8],'suitability':[1.05,2.2,1.85,1.6]}[name]
 for col,w in zip(t.columns,widths):col.width=Inches(w)
 for i,row in enumerate([heads]+rows):
  cells=t.rows[0].cells if i==0 else t.add_row().cells
  for j,(cell,value) in enumerate(zip(cells,row)):
   cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   p=cell.paragraphs[0];p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(4);p.paragraph_format.line_spacing=1.02
   r=p.add_run(value);r.font.size=Pt(9 if name=='threat' else 9.5);r.bold=(i==0)
   if name in ['sizes','benchmark'] and j>=2:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
   pr=cell._tc.get_or_add_tcPr();shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E4EAF0' if i==0 else ('F5F7F8' if i%2 else 'FFFFFF'));pr.append(shade)
   bor=OxmlElement('w:tcBorders')
   for side in ['top','left','bottom','right']:
    e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');bor.append(e)
   pr.append(bor);mar=OxmlElement('w:tcMar')
   for side in ['top','left','bottom','right']:
    e=OxmlElement('w:'+side);e.set(qn('w:w'),'70');e.set(qn('w:type'),'dxa');mar.append(e)
   pr.append(mar)
  trpr=t.rows[i]._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
  if i==0:trpr.append(OxmlElement('w:tblHeader'))
 p=doc.add_paragraph();p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(0)

n=0
for s in C['sections']:
 heading(s['heading'])
 for i,item in enumerate(s['paragraphs']):
  if isinstance(item,str):
   p=para(item)
   if i+1<len(s['paragraphs']) and isinstance(s['paragraphs'][i+1],dict) and 'equation' in s['paragraphs'][i+1]:p.paragraph_format.keep_with_next=True
  elif 'subheading' in item:heading(item['subheading'],2)
  elif 'figure' in item:figure(item['figure'])
  elif 'table' in item:table(item['table'])
  elif 'equation' in item:n+=1;equation(n)

heading('References').paragraph_format.page_break_before=True
para('Sources accessed 21 October 2026. Repository references identify the original and experimental revision.','Small')
for ref in C['references']:
 p=para('', 'Reference');p.paragraph_format.left_indent=Inches(.28);p.paragraph_format.first_line_indent=Inches(-.28)
 p.add_run(f'[{ref["n"]}] '+ref['text']+' [Online]. Available: ')
 display=ref['url']
 if 'github.com/demis1997' in display:display='github.com/demis1997/quorum-custody, pinned source'
 link(p,display,ref['url']);p.add_run('. Accessed: Oct. 21, 2026.')

heading('Appendix A  Reproducibility').paragraph_format.page_break_before=True
para('The original baseline and additional application experiments use revision 691e253ccdc2756afe215a23384eb72e2d4fde3c. No application source changes were made. The package preserves the baseline document and builder, updated research content and builder, source license, exact npm lockfiles, four SVG diagrams with compatibility PNGs, experiment scripts, and sanitized results.')
para('Environment: macOS ARM64, Apple M3 Max, Darwin 25.5.0, Node 22.20.0, npm 10.9.3. Anvil 1.3.5-nightly, commit 9cd8a9511bc7206cb251721fa42420a76c668ae3. Every chain operation used synthetic local balances. Event logs use UTC; the publication date uses Asia/Nicosia.')
para('From the research-package directory, install and run in this order. Node 22.20.0 and Anvil must be on PATH. The setup script retrieves the public source only when it is absent and refuses a different revision. Registry access is required for dependency installation. No real credentials are needed.')
for cmd in ['bash experiments/setup.sh','bash experiments/run.sh']:
 p=para(cmd);p.paragraph_format.left_indent=Inches(.15)
 for r in p.runs:r.font.name='Courier New';r.font.size=Pt(9.5)
para('The runner invokes node --import tsx --test for the repository’s test files through its installed loader, then executes security-experiments.mjs and verification-benchmark.mjs. The former starts one loopback Anvil process on port 18545 by default and terminates only that process. Set RESEARCH_PORT to an unused loopback port if necessary. It does not start Docker or modify another service.')
para('The benchmark uses five warmups and seven batches of thirty valid verifications per scheme over one 32 byte message. Keys and signatures are freshly generated and never logged. Results preserve each batch mean, exact public sizes, runtime and library versions. The hybrid row is a logical AND of two verifiers, not a new composite signature construction.')
para('For full stack work, follow the repository’s pinned native build and isolated local profile instructions, then execute npm run integration and npm run resilience. Those commands are provided for future work and were Not Executed here. They require PostgreSQL and the compiled cb-mpc adapter with patched OpenSSL. Do not interpret the research runner as a substitute for them.')
para('Document regeneration uses build_paper.py with python-docx and the supplied figure assets. toc-pages.json contains automatically derived final page locations for the cached TOC field. The distribution PDF was generated and every page inspected. Reflow after editing requires updating the TOC field and reexporting the PDF.')

heading('Appendix B  Experiment Protocols and Results').paragraph_format.page_break_before=True
for title,body in [
('A  Divergent ledger protocol','Expected outcome: valid individual approvals do not establish one budget history. Threat: A equivocates about local history, B and C retain honest partial views, and both intents already have sufficient approvals. Procedure: run the two authorize calls at nonces 0 and 1, apply the declared ledger model to A+B and A+C, then apply the central serial model. Observed: 600042000000000000 wei per reservation, 1200084000000000000 wei combined, versus a 1000000000000000000 wei limit. The central model rejects the second request. Relevant result: A record in security-experiments.json. Actual SQL reservations and native signing: Not Executed.'),
('B  Expiry protocol','Expected outcome: application decision rejects after expiry while unchanged signed bytes remain acceptable to Anvil. Threat: a holder retains previously authorized ordinary signed bytes. Procedure: create and approve a 0.01 ETH type 2 transaction, sign with a synthetic complete key, persist bytes, wait for real wall clock expiry, call validateBroadcastDecision, then submit directly. Observed: approval_expired; Anvil receipt status 1 in block 1; duplicate submission rejected with nonce too low; account nonce remains 1. Full worker and threshold signature generation: Not Executed.'),
('C  Rollback protocol','Expected outcome: an independently retained latest checkpoint rejects obsolete backup state, while joint rollback falls outside that assumption. Threat: compare a stale backup holder with a privileged actor able to write the checkpoint and use the backup password. Procedure: backup synthetic prepared session state, advance ledger/checkpoint, remove storage, attempt stale/missing/changed checkpoint restores, then restore both matching old state and checkpoint. Observed: stale_recovery_backup, ENOENT, and stale_recovery_backup respectively; joint rollback accepted. Relocation also accepted under the same administrator. Live replay, independent hosts, and actual interrupted MPC/RPC recovery: Not Executed.'),
('D  Availability protocol','Expected outcome: transition permissions admit selected retry states and forbid terminal state regression; no native availability conclusion follows. Threat: interrupted work and duplicate delivery. Procedure: call ready to signing, signing to ready, signing to blocked, signed to broadcast, and confirmed to signing; resubmit B’s included bytes. Observed: permitted transitions succeed, confirmed to signing throws invalid_transition, and the chain rejects the duplicate nonce. Real one signer outage, DKG outage, coordinator kill, timeout, participant disconnect, and native retry: Not Executed.')]:
 p=para(title);p.runs[0].bold=True;para(body)
B=next(x for x in C['results']['records'] if x['experiment']=='B')
para('B public transaction identifier','Small')
p=para(B['transactionHash'],'Small')
para('UTC event times: authorized '+B['authorizedAt']+'; approval expiry '+B['expiresAt']+'; application rejection '+B['rejectedAt']+'; submission/receipt record '+B['submittedAt']+'. Policy version 1. The local chain has been stopped; this hash is not a public network explorer reference.','Small')
para('Complete commands, all event records, benchmark batch values, and the repeated 26 test TAP log are supplied in results/. The public signed transaction file contains only local synthetic transaction bytes. No share, approval private key, password, bearer credential, or real asset is included. RESEARCH_LIMITATIONS.md enumerates every unexecuted scope.','Small')

# Populate an actual TOC field with automatically generated linked cached entries.
parent=toc_anchor._p.getparent();idx=parent.index(toc_anchor._p)
for i,h in enumerate(headings):
 p=doc.add_paragraph(style='TOC 1' if h['level']==1 else 'TOC 2')
 if i==0:
  r=p.add_run();f=OxmlElement('w:fldChar');f.set(qn('w:fldCharType'),'begin');r._r.append(f)
  r=p.add_run();instr=OxmlElement('w:instrText');instr.set(qn('xml:space'),'preserve');instr.text=' TOC \\o "1-2" \\h \\z \\u ';r._r.append(instr)
  r=p.add_run();f=OxmlElement('w:fldChar');f.set(qn('w:fldCharType'),'separate');r._r.append(f)
 link(p,h['text'],anchor=h['anchor']);p.add_run('\t'+str(pages.get(h['text'],1)))
 if i==len(headings)-1:
  r=p.add_run();f=OxmlElement('w:fldChar');f.set(qn('w:fldCharType'),'end');r._r.append(f)
 parent.insert(idx+i,p._p)
parent.remove(toc_anchor._p)
for root in [doc.styles.element,doc.element]:
 for border in list(root.iter(qn('w:pBdr'))):border.getparent().remove(border)
out=R/'paper/custody-security-research.docx';doc.save(out)
(A/'heading-manifest.json').write_text(json.dumps(headings,indent=2))
print('Saved',out,'headings',len(headings))
